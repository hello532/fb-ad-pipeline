#!/usr/bin/env python3
"""memory_scan.py — 纯 Python 内存扫描 FB Ad Library（无 HTTP / 无 FastAPI）

目标：
1. 词库整表载入内存
2. 直接跑 scrapling / playwright 抓取，不启 127.0.0.1:8000
3. 结果先落 MemoryStore，再可选 flush JSON
4. 实时吞吐 + 一天产能估算

用法：
  # 小样本测速 + 日产能估算（推荐先跑）
  python3 memory_scan.py --bench --limit 8 --workers 4

  # 从词库/队列扫 N 个关键词
  python3 memory_scan.py --library --limit 50 --workers 5

  # 指定关键词
  python3 memory_scan.py --keywords "knee brace,posture corrector" --workers 3

  # 持续跑到达到目标或被 Ctrl-C（适合估一天）
  python3 memory_scan.py --library --limit 500 --workers 5 --hours 2

环境变量（复用现有 pipeline）：
  FB_SCRAPE_WORKERS / FB_MAX_PAGES / FB_SCRAPLING_MODE / FB_DELAY_* / FB_DAILY_QUOTA
"""
from __future__ import annotations

import argparse
import json
import os
import signal
import sys
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from core.logger import setup_logging, get_logger  # noqa: E402
from pipeline.paths import LOG_DIR, RESULTS_DIR  # noqa: E402

log = get_logger("memory_scan")

# 历史实测（本机 raw JSON mtime 跨度）— 作为对照基准，不是承诺
HISTORICAL_PEAK_KPH = {
    "2026-06-16": 224.7,  # 3226 files / 14.36h
    "2026-06-17": 238.4,  # 2926 files / 12.27h
}


# ── Memory store ────────────────────────────────────────────────────────────


@dataclass
class KwResult:
    keyword: str
    ok: bool
    status: str
    ads_count: int = 0
    duration_s: float = 0.0
    source: str = ""
    file: Optional[str] = None
    error: str = ""
    ads: list[dict[str, Any]] = field(default_factory=list)


@dataclass
class MemoryStore:
    """纯内存工作集：词队列 + 抓取结果 + 广告索引 + 吞吐统计。"""

    started_at: float = field(default_factory=time.time)
    keywords: list[str] = field(default_factory=list)
    results: list[KwResult] = field(default_factory=list)
    ads_by_id: dict[str, dict[str, Any]] = field(default_factory=dict)
    ads_by_keyword: dict[str, list[str]] = field(default_factory=lambda: defaultdict(list))
    status_counts: Counter = field(default_factory=Counter)
    source_counts: Counter = field(default_factory=Counter)
    stop_requested: bool = False

    def add(self, row: KwResult) -> None:
        self.results.append(row)
        self.status_counts[row.status] += 1
        if row.source:
            self.source_counts[row.source] += 1
        for ad in row.ads:
            aid = str(
                ad.get("ad_archive_id")
                or ad.get("ad_id")
                or ad.get("id")
                or ""
            ).strip()
            if not aid:
                # 无 id 时用内容 hash 兜底，避免丢广告
                aid = f"anon:{hash((ad.get('page_name'), ad.get('link_url'), ad.get('body'))) & 0xFFFFFFFF:x}"
            if aid not in self.ads_by_id:
                self.ads_by_id[aid] = ad
            self.ads_by_keyword[row.keyword].append(aid)

    @property
    def elapsed_s(self) -> float:
        return max(0.001, time.time() - self.started_at)

    @property
    def done(self) -> int:
        return len(self.results)

    @property
    def ok_count(self) -> int:
        return sum(1 for r in self.results if r.ok)

    @property
    def ads_total(self) -> int:
        return len(self.ads_by_id)

    @property
    def kw_per_hour(self) -> float:
        return self.done * 3600.0 / self.elapsed_s

    @property
    def ads_per_hour(self) -> float:
        return self.ads_total * 3600.0 / self.elapsed_s

    def snapshot(self) -> dict[str, Any]:
        dur = [r.duration_s for r in self.results if r.duration_s > 0]
        avg_dur = sum(dur) / len(dur) if dur else 0.0
        ok_ads = [r.ads_count for r in self.results if r.ok]
        avg_ads = sum(ok_ads) / len(ok_ads) if ok_ads else 0.0
        kph = self.kw_per_hour
        return {
            "ok": True,
            "mode": "memory",
            "http": False,
            "started_at": datetime.fromtimestamp(self.started_at).isoformat(timespec="seconds"),
            "elapsed_s": round(self.elapsed_s, 2),
            "requested": len(self.keywords),
            "done": self.done,
            "ok_count": self.ok_count,
            "unique_ads": self.ads_total,
            "status": dict(self.status_counts),
            "sources": dict(self.source_counts),
            "avg_duration_s": round(avg_dur, 2),
            "avg_ads_per_ok_kw": round(avg_ads, 2),
            "throughput": {
                "kw_per_hour": round(kph, 1),
                "ads_per_hour": round(self.ads_per_hour, 1),
            },
            "daily_capacity": _estimate_daily(kph, avg_ads),
            "historical_peak_kph": HISTORICAL_PEAK_KPH,
        }


def _estimate_daily(kph: float, avg_ads: float) -> dict[str, Any]:
    """按当前实测速率外推一天产能。含可持续档（不被短 bench 虚高带跑）。"""
    if kph <= 0:
        return {"note": "尚无有效样本"}
    # 现实：去掉静默/风控/吃饭，按 14h 有效
    # 激进：24h 无配额（需 FB_DAILY_QUOTA=0 且风控平稳）
    # 保守：8h 工作窗 + 20% 风控折损
    conservative_h = 8 * 0.8
    realistic_h = 14
    aggressive_h = 24
    # 可持续上限：取 min(实测, 历史峰值*1.3, 800) — 防止 L0 短 burst 外推到不现实
    hist_peak = max(HISTORICAL_PEAK_KPH.values()) if HISTORICAL_PEAK_KPH else 240.0
    sustainable_kph = min(kph, hist_peak * 1.3, 800.0)
    if kph > sustainable_kph * 1.5:
        burst_note = (
            f"本轮 {kph:.0f} kw/h 多为 L0-direct 短 burst；"
            f"可持续按 {sustainable_kph:.0f} kw/h 估，长跑会掉到历史 225-240 附近"
        )
    else:
        burst_note = "本轮速率接近可持续区间"
    return {
        "based_on_kph": round(kph, 1),
        "sustainable_kph": round(sustainable_kph, 1),
        "avg_ads_per_ok_kw": round(avg_ads, 2),
        "burst_from_this_run": {
            "conservative_8h_x0.8": {
                "keywords": int(kph * conservative_h),
                "ads": int(kph * conservative_h * avg_ads),
                "hours": conservative_h,
            },
            "realistic_14h": {
                "keywords": int(kph * realistic_h),
                "ads": int(kph * realistic_h * avg_ads),
                "hours": realistic_h,
            },
            "aggressive_24h": {
                "keywords": int(kph * aggressive_h),
                "ads": int(kph * aggressive_h * avg_ads),
                "hours": aggressive_h,
            },
        },
        "sustainable_day": {
            "conservative_8h_x0.8": {
                "keywords": int(sustainable_kph * conservative_h),
                "ads": int(sustainable_kph * conservative_h * avg_ads),
                "hours": conservative_h,
            },
            "realistic_14h": {
                "keywords": int(sustainable_kph * realistic_h),
                "ads": int(sustainable_kph * realistic_h * avg_ads),
                "hours": realistic_h,
            },
            "aggressive_24h": {
                "keywords": int(sustainable_kph * aggressive_h),
                "ads": int(sustainable_kph * aggressive_h * avg_ads),
                "hours": aggressive_h,
            },
        },
        "recommended_daily_keywords": int(sustainable_kph * realistic_h),
        "note": burst_note,
    }



# ── Keyword loading ─────────────────────────────────────────────────────────


def _parse_text_items(text: str) -> list[str]:
    import re

    seen: set[str] = set()
    out: list[str] = []
    for part in re.split(r"[\n,;\t]+", text or ""):
        item = part.strip().strip("\"'").lower()
        if 2 <= len(item) <= 120 and item not in seen:
            seen.add(item)
            out.append(item)
    return out


def load_keywords_from_library(limit: int) -> list[str]:
    """优先 scan_queue_loader（DB 未扫词）；失败则回退 scan.txt。"""
    try:
        from pipeline.scan_instructions import scan_queue_loader

        rows = scan_queue_loader()
        kws = [r["keyword"] for r in rows if r.get("keyword")]
        if kws:
            return kws[: max(1, limit)]
    except Exception as exc:
        log.warning(f"scan_queue_loader 失败, 回退 scan.txt: {exc}")

    from pipeline.paths import SCAN_FILE

    kws: list[str] = []
    seen: set[str] = set()
    if SCAN_FILE.exists():
        for line in SCAN_FILE.read_text(encoding="utf-8", errors="ignore").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            kw = line.split("|")[0].strip().lower()
            if 2 <= len(kw) <= 120 and kw not in seen:
                seen.add(kw)
                kws.append(kw)
            if len(kws) >= limit:
                break
    return kws[: max(1, limit)]


def load_keywords_from_rabbit_hole(
    limit: int,
    *,
    seeds: Optional[list[str]] = None,
    expand: bool = True,
    max_per_layer: int = 8,
) -> tuple[list[str], dict[str, Any]]:
    """兔子洞选词：队列文件 + 可选 10 层 expand。

    返回 (keywords, meta)。全程内存，不经 HTTP。
    """
    from pipeline.paths import RABBIT_QUEUE

    meta: dict[str, Any] = {
        "method": "rabbit_hole",
        "queue_file": str(RABBIT_QUEUE),
        "queue_exists": RABBIT_QUEUE.exists(),
        "expand": expand,
        "seeds": [],
        "from_queue": 0,
        "from_expand": 0,
    }
    seen: set[str] = set()
    out: list[str] = []

    def _add(kw: str) -> bool:
        kw = (kw or "").strip().lower()
        if not (2 <= len(kw) <= 120) or kw in seen:
            return False
        # 跳过明显噪声/非选品词
        junk = ("free shipping", "learn more", "click here", "over 100", "over 50", "over 1", "wobi")
        if kw in junk or kw.startswith("http"):
            return False
        seen.add(kw)
        out.append(kw)
        return True

    # 1) 现有兔子洞队列
    if RABBIT_QUEUE.exists():
        for line in RABBIT_QUEUE.read_text(encoding="utf-8", errors="ignore").splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            kw = line.split("|")[0].strip()
            if _add(kw):
                meta["from_queue"] += 1
            if len(out) >= limit:
                break

    # 2) 种子 expand（真正的兔子洞 10 层选品法）
    seed_list = list(seeds or [])
    if not seed_list:
        # 队列前几个当种子再扩一轮
        seed_list = out[: min(5, len(out))] or [
            "knee pain",
            "neck massager",
            "posture corrector",
            "plantar fasciitis",
            "sciatica",
        ]
    meta["seeds"] = seed_list

    if expand and len(out) < limit:
        try:
            from pipeline.keyword_evolution import rabbit_hole_expand, _flatten_expansion

            for seed in seed_list:
                expanded = rabbit_hole_expand(seed)
                for item in _flatten_expansion(expanded, max_per_layer=max_per_layer):
                    if _add(item.get("keyword") or ""):
                        meta["from_expand"] += 1
                    if len(out) >= limit:
                        break
                if len(out) >= limit:
                    break
        except Exception as exc:
            log.warning(f"rabbit_hole_expand 失败: {exc}")
            meta["expand_error"] = str(exc)

    meta["selected"] = len(out[:limit])
    return out[: max(1, limit)], meta


# ── Scrape one keyword into memory ──────────────────────────────────────────


def _extract_ads_from_path(path: Path) -> tuple[list[dict[str, Any]], str]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return [], ""
    ads = data.get("results") or data.get("ads") or []
    if not isinstance(ads, list):
        ads = []
    source = str(data.get("source") or "")
    return ads, source


def _clean_text(value: Any) -> str:
    """清理 FB JSON 转义/控制字符，避免 Excel 乱码。"""
    if value is None:
        return ""
    s = str(value)
    # JSON 风格转义
    s = s.replace("\\/", "/")
    s = s.replace("\\n", "\n").replace("\\r", "\r").replace("\\t", "\t")
    s = s.replace('\\"', '"').replace("\\'", "'")
    # \uXXXX
    try:
        import re as _re

        s = _re.sub(r"\\u([0-9a-fA-F]{4})", lambda m: chr(int(m.group(1), 16)), s)
    except Exception:
        pass
    # 去控制符 / surrogate
    s = "".join(ch for ch in s if ch >= " " or ch in "\n\t")
    s = s.encode("utf-8", "ignore").decode("utf-8", "ignore")
    return s.strip()


def _clean_url(url: Any) -> str:
    s = _clean_text(url)
    if not s:
        return ""
    s = s.replace("\\/", "/").strip().strip('"').strip("'")
    # 有些字段被双转义
    while "\\/" in s:
        s = s.replace("\\/", "/")
    if s.startswith("//"):
        s = "https:" + s
    return s


def _to_iso_date(value: Any) -> str:
    if value in (None, "", 0, "0"):
        return ""
    if isinstance(value, str) and len(value) >= 8 and value[4] == "-":
        return value[:10]
    try:
        iv = int(float(value))
        if iv > 1_000_000_000_000:  # ms
            iv //= 1000
        if iv > 1_000_000:
            return datetime.utcfromtimestamp(iv).strftime("%Y-%m-%d")
    except Exception:
        pass
    return _clean_text(value)[:10]


def _active_days(start_iso: str) -> int:
    if not start_iso:
        return 0
    try:
        start = datetime.strptime(start_iso[:10], "%Y-%m-%d")
        return max(0, (datetime.utcnow() - start).days)
    except Exception:
        return 0


def _ad_fields(ad: dict[str, Any]) -> dict[str, Any]:
    """广告字段归一 + 清洗（兼容 DOM / GraphQL）— 保留运营级富字段。"""
    from pipeline.domain import extract_domain_from_url

    link = _clean_url(
        ad.get("link_url")
        or ad.get("linkUrl")
        or ad.get("url")
        or ad.get("snapshot_url")
        or ""
    )
    cta = _clean_text(
        ad.get("cta_type") or ad.get("cta_text") or ad.get("cta") or ad.get("call_to_action_type") or ""
    ).upper()
    page = _clean_text(ad.get("page_name") or ad.get("pageName") or ad.get("page") or "")
    title = _clean_text(ad.get("title") or ad.get("link_caption") or "")
    body = _clean_text(ad.get("body") or ad.get("ad_text") or title or "")
    domain = _clean_text(ad.get("domain") or "")
    if not domain and link:
        domain = extract_domain_from_url(link)
    if (not domain or domain.endswith("facebook.com") or domain.endswith("instagram.com")) and page:
        from pipeline.domain import KNOWN_DOMAINS

        if page in KNOWN_DOMAINS:
            domain = KNOWN_DOMAINS[page]
    domain = domain.lower().lstrip("www.") if domain else ""

    start_raw = ad.get("start_date_iso") or ad.get("start_date") or ad.get("startDate")
    start_iso = _to_iso_date(start_raw)
    end_iso = _to_iso_date(ad.get("end_date") or ad.get("end_date_iso") or "")
    likes = ad.get("page_like_count") or ad.get("page_likes") or 0
    try:
        likes = int(likes or 0)
    except Exception:
        likes = 0

    platforms = ad.get("publisher_platforms") or ad.get("platforms") or []
    if isinstance(platforms, str):
        platforms = [p.strip() for p in platforms.split(",") if p.strip()]
    if not isinstance(platforms, list):
        platforms = []

    aid = str(ad.get("ad_archive_id") or ad.get("ad_id") or ad.get("id") or "").strip()
    out = {
        "ad_archive_id": aid,
        "page_name": page,
        "page_id": str(ad.get("page_id") or ""),
        "page_like_count": likes,
        "link_url": link,
        "cta": cta,
        "domain": domain,
        "title": title[:200],
        "body": body[:800],
        "start_date": start_iso,
        "end_date": end_iso,
        "active_days": _active_days(start_iso),
        "is_active": bool(ad.get("is_active", ad.get("isActive", True))),
        "display_format": _clean_text(ad.get("display_format") or ""),
        "publisher_platforms": [str(p) for p in platforms],
        "spend": ad.get("spend") if isinstance(ad.get("spend"), dict) else ad.get("spend"),
        "impressions": ad.get("impressions") if isinstance(ad.get("impressions"), dict) else ad.get("impressions"),
        "currency": _clean_text(ad.get("currency") or ""),
        "collation_count": ad.get("collation_count") or 1,
        "ig_handle": _clean_text(ad.get("ig_handle") or ad.get("instagram_actor_name") or ""),
        "ig_followers": ad.get("ig_followers") or 0,
        "image_url": _clean_url(ad.get("image_url") or ad.get("original_image_url") or ""),
        "video_url": _clean_url(ad.get("video_url") or ad.get("video_hd_url") or ""),
        "carousel_images": ad.get("carousel_images") or [],
        "carousel_videos": ad.get("carousel_videos") or [],
        "ad_snapshot_url": _clean_url(ad.get("ad_snapshot_url") or ad.get("snapshot_url") or ""),
        "link_caption": _clean_text(ad.get("link_caption") or "")[:160],
        "link_description": _clean_text(ad.get("link_description") or "")[:200],
        "page_categories": ad.get("page_categories") or [],
        "byline": _clean_text(ad.get("byline") or ""),
    }
    try:
        from pipeline.extreme_table import enrich_ad

        out = enrich_ad(out)
    except Exception:
        pass
    return out





def scrape_one_pure_memory(
    keyword: str,
    *,
    max_scrolls: int,
    headless: bool,
    country: str,
    max_pages: Optional[int],
    keep_ads: bool,
) -> KwResult:
    """纯内存抓取：走 scrapling L0/L1，不写 raw JSON 文件。"""
    from pipeline.commerce_filter import filter_commerce_ads
    from pipeline.proxy_pool import scrape_proxy_context
    from pipeline.scrapling_scraper import (
        L0_DIRECT_SAVE_MIN_ADS,
        MAX_PAGES,
        quick_scan,
        stealthy_scan,
        _patchright_paginate,
    )

    t0 = time.time()
    pages = max_pages if max_pages is not None else MAX_PAGES
    try:
        with scrape_proxy_context():
            l0 = quick_scan(keyword, country)
            l0_down = l0 is None
            if l0_down:
                # L0 直连不可用 ≠ 词没广告：落空对象，让后续 L1/L1.5 兜底路径接管
                # （此前直接 return，导致 L1 明明可用也整批 l0_unavailable）
                l0 = {"count": 0, "ads": []}
            if not l0_down and l0.get("challenge"):
                return KwResult(
                    keyword=keyword,
                    ok=False,
                    status="challenge",
                    duration_s=round(time.time() - t0, 2),
                    source="L0-fetcher",
                )
            if not l0_down and int(l0.get("count") or 0) == 0:
                return KwResult(
                    keyword=keyword,
                    ok=False,
                    status="empty",
                    duration_s=round(time.time() - t0, 2),
                    source="L0-fetcher",
                )

            result = None
            source = "L0-fetcher"

            if pages > 1:
                l15 = _patchright_paginate(
                    keyword, country, headless=headless, max_pages=pages
                )
                if l15 and l15.get("challenge"):
                    return KwResult(
                        keyword=keyword,
                        ok=False,
                        status="challenge",
                        duration_s=round(time.time() - t0, 2),
                        source="L1.5-patchright",
                    )
                if l15 and int(l15.get("count") or 0) > 0:
                    result = l15
                    source = str(l15.get("source") or "L1.5-patchright")

            if result is None and pages <= 1 and L0_DIRECT_SAVE_MIN_ADS and int(l0["count"]) >= L0_DIRECT_SAVE_MIN_ADS:
                ads, _drop = filter_commerce_ads(list(l0.get("ads") or []))
                result = {"ads": ads, "count": len(ads)}
                source = "L0-fetcher-direct-memory"

            if result is None:
                l1 = stealthy_scan(
                    keyword,
                    country,
                    headless=headless,
                    max_scrolls=max(1, min(max_scrolls, 8)),
                    max_pages=1,
                )
                if l1 and l1.get("challenge"):
                    return KwResult(
                        keyword=keyword,
                        ok=False,
                        status="challenge",
                        duration_s=round(time.time() - t0, 2),
                        source="L1-stealthy",
                    )
                if l1 and int(l1.get("count") or 0) > 0:
                    result = l1
                    source = str(l1.get("source") or "L1-stealthy")
                    if l0_down:
                        source += "+L0-down-fallback"
                elif l1 and int(l1.get("count") or 0) == 0:
                    return KwResult(
                        keyword=keyword,
                        ok=False,
                        status="empty",
                        duration_s=round(time.time() - t0, 2),
                        source="L1-stealthy",
                    )

            if result is None:
                # L0 有命中但未达 direct 阈值且 L1 失败：仍用 L0（内存兜底）
                ads, _drop = filter_commerce_ads(list(l0.get("ads") or []))
                if not ads:
                    if l0_down:
                        # 双层都失败才报 unavailable，并给可执行恢复路径
                        return KwResult(
                            keyword=keyword,
                            ok=False,
                            status="l0_unavailable",
                            duration_s=round(time.time() - t0, 2),
                            source="L0+L1-failed",
                            error="L0直连与L1均无结果：检查代理池(scrape_proxy_context)/FB登录态；"
                                  "或改走 fb-scan-pipeline-v8.py --v3-scrape '词' --scrape-headful",
                        )
                    return KwResult(
                        keyword=keyword,
                        ok=False,
                        status="empty_or_failed",
                        duration_s=round(time.time() - t0, 2),
                        source="L0-fallback",
                    )
                result = {"ads": ads, "count": len(ads)}
                source = "L0-fetcher-memory-fallback"

        ads_in = list(result.get("ads") or [])
        # L1/L1.5 内部可能已 filter；再归一化字段
        norm = [_ad_fields(a) if "raw" not in a else a for a in ads_in]
        # 若还是 raw，统一包一层
        final_ads = []
        for a in ads_in:
            final_ads.append(_ad_fields(a))

        dur = round(time.time() - t0, 2)
        return KwResult(
            keyword=keyword,
            ok=True,
            status="memory",
            ads_count=len(final_ads),
            duration_s=dur,
            source=source,
            file=None,
            ads=final_ads if keep_ads else [],
        )
    except Exception as exc:
        log.exception(f"pure memory scrape 异常: {keyword}")
        return KwResult(
            keyword=keyword,
            ok=False,
            status="error",
            duration_s=round(time.time() - t0, 2),
            error=str(exc),
        )


def scrape_one_to_memory(
    keyword: str,
    *,
    max_scrolls: int,
    headless: bool,
    country: str,
    max_pages: Optional[int],
    keep_ads: bool,
    pure_memory: bool = True,
) -> KwResult:
    """抓一个词 → KwResult。默认纯内存；pure_memory=False 时走旧 scrape_keyword（会写 raw）。"""
    if pure_memory:
        return scrape_one_pure_memory(
            keyword,
            max_scrolls=max_scrolls,
            headless=headless,
            country=country,
            max_pages=max_pages,
            keep_ads=keep_ads,
        )

    from pipeline.proxy_pool import scrape_proxy_context
    from pipeline.scraper import scrape_keyword

    t0 = time.time()
    try:
        with scrape_proxy_context():
            out = scrape_keyword(
                keyword,
                max_scrolls=max_scrolls,
                headless=headless,
                country=country,
                max_pages=max_pages,
            )
        dur = round(time.time() - t0, 2)
        if out == "CHALLENGE":
            return KwResult(keyword=keyword, ok=False, status="challenge", duration_s=dur)
        if out and isinstance(out, Path):
            ads, source = _extract_ads_from_path(out)
            return KwResult(
                keyword=keyword,
                ok=True,
                status="saved",
                ads_count=len(ads),
                duration_s=dur,
                source=source,
                file=str(out),
                ads=[_ad_fields(a) for a in ads] if keep_ads else [],
            )
        return KwResult(
            keyword=keyword,
            ok=False,
            status="empty_or_failed",
            duration_s=dur,
        )
    except Exception as exc:
        log.exception(f"memory scrape 异常: {keyword}")
        return KwResult(
            keyword=keyword,
            ok=False,
            status="error",
            duration_s=round(time.time() - t0, 2),
            error=str(exc),
        )



# ── Runner ──────────────────────────────────────────────────────────────────


def run_memory_scan(
    keywords: list[str],
    *,
    workers: int = 4,
    max_scrolls: int = 5,
    headless: bool = True,
    country: str = "US",
    max_pages: Optional[int] = None,
    delay_sec: float = 0.1,
    keep_ads: bool = True,
    hours: float = 0.0,
    progress_every: int = 5,
    pure_memory: bool = True,
) -> MemoryStore:
    from concurrent.futures import wait, FIRST_COMPLETED

    store = MemoryStore(keywords=list(keywords))
    workers = max(1, min(workers, len(keywords) or 1))
    deadline = (time.time() + hours * 3600.0) if hours and hours > 0 else None

    def _on_sig(_signum, _frame):
        store.stop_requested = True
        log.warning("收到停止信号，当前批次结束后退出")

    prev_int = signal.signal(signal.SIGINT, _on_sig)
    prev_term = signal.signal(signal.SIGTERM, _on_sig)

    log.info(
        f"memory_scan start: n={len(keywords)} workers={workers} "
        f"country={country} pure_memory={pure_memory} keep_ads={keep_ads} hours={hours or 'inf'}"
    )

    try:
        with ThreadPoolExecutor(max_workers=workers) as ex:
            it = iter(keywords)
            futures: dict[Any, str] = {}

            def _fill() -> None:
                while len(futures) < workers and not store.stop_requested:
                    if deadline and time.time() >= deadline:
                        store.stop_requested = True
                        return
                    try:
                        kw = next(it)
                    except StopIteration:
                        return
                    fut = ex.submit(
                        scrape_one_to_memory,
                        kw,
                        max_scrolls=max_scrolls,
                        headless=headless,
                        country=country,
                        max_pages=max_pages,
                        keep_ads=keep_ads,
                        pure_memory=pure_memory,
                    )
                    futures[fut] = kw

            _fill()
            while futures:
                done_set, _ = wait(futures.keys(), return_when=FIRST_COMPLETED)
                for fut in done_set:
                    kw = futures.pop(fut)
                    try:
                        row = fut.result()
                    except Exception as exc:
                        row = KwResult(
                            keyword=kw,
                            ok=False,
                            status="error",
                            error=str(exc),
                        )
                    store.add(row)
                    n = store.done
                    if progress_every > 0 and (n % progress_every == 0 or n >= len(keywords)):
                        snap = store.snapshot()["throughput"]
                        log.info(
                            f"progress {n}/{len(keywords)} ok={store.ok_count} "
                            f"ads={store.ads_total} kph={snap['kw_per_hour']} "
                            f"last={row.keyword}:{row.status}/{row.ads_count}"
                        )
                    if delay_sec > 0:
                        time.sleep(min(delay_sec, 1.0))
                if not store.stop_requested:
                    _fill()
    finally:
        signal.signal(signal.SIGINT, prev_int)
        signal.signal(signal.SIGTERM, prev_term)

    return store


def _commerce_score(ad: dict[str, Any]) -> tuple[int, list[str]]:
    """单条广告赚钱信号分（不是利润结论）。"""
    score = 0
    reasons: list[str] = []
    cta = (ad.get("cta") or "").upper()
    url = (ad.get("link_url") or "").lower()
    domain = (ad.get("domain") or "").lower()

    if cta in {"SHOP_NOW", "BUY_NOW", "ORDER_NOW", "GET_OFFER", "ADD_TO_CART", "VIEW_PRODUCT"}:
        score += 25
        reasons.append(f"CTA={cta}")
    elif cta in {"LEARN_MORE", "GET_PROMOTIONS", "SUBSCRIBE"}:
        score += 8
        reasons.append(f"CTA弱={cta}")

    if any(x in url for x in ("/products/", "/product/", "/collections/", "myshopify.com", "/cart")):
        score += 25
        reasons.append("电商路径")
    if "myshopify.com" in domain or domain.endswith(".shop"):
        score += 10
        reasons.append("独立站形态")

    days = int(ad.get("active_days") or 0)
    if days >= 90:
        score += 20
        reasons.append(f"长寿{days}天")
    elif days >= 30:
        score += 12
        reasons.append(f"验证期{days}天")
    elif days >= 7:
        score += 5
        reasons.append(f"测试期{days}天")

    plat = ("amazon.", "temu.", "shein.", "walmart.", "ebay.", "etsy.", "aliexpress.")
    if any(p in domain for p in plat) or any(p in url for p in plat):
        score -= 40
        reasons.append("平台货")

    likes = int(ad.get("page_like_count") or 0)
    if 1000 <= likes <= 800_000:
        score += 8
        reasons.append(f"粉丝{likes}")
    return max(0, min(100, score)), reasons


def build_money_tables(store: MemoryStore) -> dict[str, list[list[Any]]]:
    """竞争对手 / 赚钱产品 / 清洗广告明细。"""
    aid_kws: dict[str, set[str]] = defaultdict(set)
    for kw, aids in store.ads_by_keyword.items():
        for aid in aids:
            aid_kws[aid].add(kw)

    def _norm(ad0: dict[str, Any]) -> dict[str, Any]:
        ad = dict(ad0)
        ad["link_url"] = _clean_url(ad.get("link_url"))
        ad["page_name"] = _clean_text(ad.get("page_name"))
        ad["body"] = _clean_text(ad.get("body"))
        ad["title"] = _clean_text(ad.get("title"))
        ad["cta"] = _clean_text(ad.get("cta")).upper()
        ad["domain"] = _clean_text(ad.get("domain")).lower().lstrip("www.")
        if not ad["domain"] and ad["link_url"]:
            from pipeline.domain import extract_domain_from_url

            ad["domain"] = extract_domain_from_url(ad["link_url"])
        ad["start_date"] = _to_iso_date(ad.get("start_date"))
        ad["active_days"] = int(ad.get("active_days") or _active_days(ad["start_date"]))
        try:
            ad["page_like_count"] = int(ad.get("page_like_count") or 0)
        except Exception:
            ad["page_like_count"] = 0
        return ad

    # 竞品按 domain
    comp: dict[str, dict[str, Any]] = {}
    for aid, ad0 in store.ads_by_id.items():
        ad = _norm(ad0 if "link_url" in ad0 else _ad_fields(ad0))
        key = ad["domain"] or f"page:{(ad.get('page_name') or 'unknown').lower()}"
        row = comp.setdefault(
            key,
            {
                "domain": ad["domain"],
                "page_names": set(),
                "ads": 0,
                "landings": set(),
                "keywords": set(),
                "ctas": Counter(),
                "max_days": 0,
                "min_start": "",
                "score_sum": 0,
                "reasons": Counter(),
                "sample_url": "",
                "likes": 0,
            },
        )
        row["ads"] += 1
        if ad.get("page_name"):
            row["page_names"].add(ad["page_name"])
        if ad.get("link_url"):
            row["landings"].add(ad["link_url"])
            if not row["sample_url"]:
                row["sample_url"] = ad["link_url"]
        row["keywords"] |= aid_kws.get(aid, set())
        if ad.get("cta"):
            row["ctas"][ad["cta"]] += 1
        row["max_days"] = max(row["max_days"], ad["active_days"])
        sd = ad.get("start_date") or ""
        if sd and (not row["min_start"] or sd < row["min_start"]):
            row["min_start"] = sd
        sc, reasons = _commerce_score(ad)
        row["score_sum"] += sc
        for r in reasons:
            row["reasons"][r] += 1
        row["likes"] = max(row["likes"], ad["page_like_count"])

    competitors_rows: list[list[Any]] = [[
        "优先级", "域名", "品牌页", "广告数", "落地页数", "覆盖关键词数",
        "主CTA", "最长投放天", "最早投放", "粉丝", "赚钱信号分", "信号", "代表落地页", "建议",
    ]]
    ranked_comp = []
    for _key, row in comp.items():
        avg = int(row["score_sum"] / max(1, row["ads"]))
        bonus = min(20, row["ads"] * 2) + min(15, len(row["keywords"]) * 3)
        final = min(100, avg + bonus // 2)
        top_cta = row["ctas"].most_common(1)[0][0] if row["ctas"] else ""
        signals = ", ".join(r for r, _ in row["reasons"].most_common(4))
        if final >= 70:
            advice = "重点跟"
        elif final >= 50:
            advice = "观察"
        elif final >= 30:
            advice = "低优先"
        else:
            advice = "跳过"
        ranked_comp.append((final, row, top_cta, signals, advice))
    ranked_comp.sort(key=lambda x: (-x[0], -x[1]["ads"]))
    for i, (final, row, top_cta, signals, advice) in enumerate(ranked_comp, 1):
        competitors_rows.append([
            i, row["domain"], " / ".join(sorted(row["page_names"])[:3]),
            row["ads"], len(row["landings"]), len(row["keywords"]), top_cta,
            row["max_days"], row["min_start"], row["likes"], final, signals,
            row["sample_url"], advice,
        ])

    # 产品按 domain+path
    products: dict[str, dict[str, Any]] = {}
    for aid, ad0 in store.ads_by_id.items():
        ad = _norm(ad0 if "link_url" in ad0 else _ad_fields(ad0))
        url = ad["link_url"]
        domain = ad["domain"]
        path_key = ""
        try:
            from urllib.parse import urlparse

            path_key = "/".join([p for p in (urlparse(url).path or "").split("/") if p][:3])
        except Exception:
            path_key = ""
        pkey = f"{domain}|{path_key}" if domain else f"page:{ad.get('page_name') or 'x'}"
        prow = products.setdefault(
            pkey,
            {
                "domain": domain,
                "path": path_key,
                "page": ad.get("page_name") or "",
                "url": url,
                "ads": 0,
                "keywords": set(),
                "ctas": Counter(),
                "max_days": 0,
                "score_best": 0,
                "reasons": Counter(),
                "title": (ad.get("title") or ad.get("body") or "")[:120],
            },
        )
        prow["ads"] += 1
        prow["keywords"] |= aid_kws.get(aid, set())
        if ad.get("cta"):
            prow["ctas"][ad["cta"]] += 1
        prow["max_days"] = max(prow["max_days"], ad["active_days"])
        sc, reasons = _commerce_score(ad)
        prow["score_best"] = max(prow["score_best"], sc)
        for r in reasons:
            prow["reasons"][r] += 1
        if url and not prow["url"]:
            prow["url"] = url

    product_rows: list[list[Any]] = [[
        "排名", "产品/落地", "域名", "品牌", "广告创意数", "覆盖词数",
        "主CTA", "最长投放天", "信号分", "赚钱信号", "落地页", "建议",
    ]]
    ranked_prod = []
    for _pk, prow in products.items():
        bonus = min(15, prow["ads"] * 2) + min(10, len(prow["keywords"]) * 2)
        final = min(100, prow["score_best"] + bonus)
        top_cta = prow["ctas"].most_common(1)[0][0] if prow["ctas"] else ""
        signals = ", ".join(r for r, _ in prow["reasons"].most_common(4))
        label = prow["path"] or prow["title"] or prow["page"] or _pk
        if final >= 70:
            advice = "可深挖/跟品验证"
        elif final >= 50:
            advice = "值得看落地页"
        elif final >= 30:
            advice = "弱信号"
        else:
            advice = "不像卖货"
        ranked_prod.append((final, prow, label, top_cta, signals, advice))
    ranked_prod.sort(key=lambda x: (-x[0], -x[1]["ads"]))
    for i, (final, prow, label, top_cta, signals, advice) in enumerate(ranked_prod[:300], 1):
        product_rows.append([
            i, str(label)[:80], prow["domain"], prow["page"], prow["ads"],
            len(prow["keywords"]), top_cta, prow["max_days"], final, signals,
            prow["url"], advice,
        ])

    ad_rows: list[list[Any]] = [[
        "ad_id", "关键词", "品牌", "域名", "CTA", "投放天", "开始日",
        "信号分", "落地页", "文案摘要",
    ]]
    for aid, ad0 in store.ads_by_id.items():
        ad = _norm(ad0 if "link_url" in ad0 else _ad_fields(ad0))
        sc, _ = _commerce_score(ad)
        ad_rows.append([
            aid,
            ", ".join(sorted(aid_kws.get(aid, set()))[:6]),
            ad.get("page_name") or "",
            ad.get("domain") or "",
            ad.get("cta") or "",
            ad.get("active_days") or 0,
            ad.get("start_date") or "",
            sc,
            ad.get("link_url") or "",
            (ad.get("body") or ad.get("title") or "")[:180],
        ])

    return {"competitors": competitors_rows, "products": product_rows, "ads": ad_rows}


def verify_landings(
    store: MemoryStore,
    *,
    limit: int = 40,
    workers: int = 6,
) -> dict[str, dict[str, Any]]:
    """对内存里的高优落地页做技术栈/卖货验证（lp_prober，不经 HTTP 服务）。"""
    from concurrent.futures import ThreadPoolExecutor, as_completed

    from pipeline.lp_prober import probe_one

    # 收集 domain -> best url（优先 /products/）
    candidates: dict[str, dict[str, Any]] = {}
    for aid, ad0 in store.ads_by_id.items():
        ad = ad0 if isinstance(ad0, dict) else {}
        url = _clean_url(ad.get("link_url"))
        domain = _clean_text(ad.get("domain")).lower().lstrip("www.")
        if not domain and url:
            from pipeline.domain import extract_domain_from_url

            domain = extract_domain_from_url(url)
        if not domain or not url:
            continue
        if not url.startswith("http"):
            continue
        # 跳过明显平台
        if any(p in domain for p in ("amazon.", "temu.", "shein.", "facebook.", "instagram.", "shopee.")):
            continue
        sc, _ = _commerce_score({
            "cta": _clean_text(ad.get("cta")).upper(),
            "link_url": url,
            "domain": domain,
            "active_days": int(ad.get("active_days") or _active_days(_to_iso_date(ad.get("start_date")))),
            "page_like_count": ad.get("page_like_count") or 0,
        })
        prev = candidates.get(domain)
        prefer = ("/products/" in url) or ("/product/" in url) or ("myshopify.com" in url)
        prev_prefer = bool(prev and (("/products/" in (prev.get("url") or "")) or ("/product/" in (prev.get("url") or ""))))
        if prev is None or sc > prev["score"] or (prefer and not prev_prefer):
            candidates[domain] = {
                "domain": domain,
                "url": url,
                "score": sc,
                "page_name": _clean_text(ad.get("page_name")),
            }

    ranked = sorted(candidates.values(), key=lambda r: -int(r["score"]))[: max(1, limit)]
    log.info(f"落地验证: 计划探测 {len(ranked)} 个域名 workers={workers}")

    out: dict[str, dict[str, Any]] = {}

    def _verdict(res: dict[str, Any]) -> str:
        err = res.get("error")
        status = int(res.get("http_status") or 0)
        tags = res.get("tags") or []
        if err or status and status >= 400 or status == 0:
            return "FAIL"
        if res.get("has_dropship"):
            return "FAIL_DROPSHIP"
        strong = any(t in tags for t in ("lp:shopify", "lp:klaviyo", "lp:subscription", "lp:sms"))
        dual = ("lp:payment" in tags and ("lp:meta_pixel" in tags or "lp:tiktok_pixel" in tags or "lp:ga4" in tags))
        if strong or dual:
            if int(res.get("reddit_bonus") or 0) >= 5 or res.get("is_single_product"):
                return "PASS_STRONG"
            return "PASS"
        if status == 200:
            return "WEAK"
        return "FAIL"

    def _run(item: dict[str, Any]) -> dict[str, Any]:
        try:
            res = probe_one(item["domain"], item["url"], conn=None)
        except Exception as exc:
            res = {
                "domain": item["domain"],
                "url": item["url"],
                "tags": [],
                "bonus": 0,
                "price_hint": None,
                "has_dropship": False,
                "http_status": 0,
                "error": str(exc)[:200],
                "reddit_bonus": 0,
                "trust_badges": 0,
                "is_single_product": False,
                "discount_stack_count": 0,
                "return_policy_days": 0,
            }
        res["page_name"] = item.get("page_name") or ""
        res["pre_score"] = item.get("score") or 0
        res["verdict"] = _verdict(res)
        # 综合验证分
        vscore = int(res.get("bonus") or 0) + int(res.get("reddit_bonus") or 0)
        if res["verdict"].startswith("PASS"):
            vscore += 20
        if res["verdict"] == "PASS_STRONG":
            vscore += 15
        if res["verdict"].startswith("FAIL"):
            vscore = min(vscore, 15)
        res["verify_score"] = max(0, min(100, vscore + int(item.get("score") or 0) // 5))
        return res

    with ThreadPoolExecutor(max_workers=max(1, min(workers, 10))) as ex:
        futs = {ex.submit(_run, item): item["domain"] for item in ranked}
        for fut in as_completed(futs):
            dom = futs[fut]
            try:
                out[dom] = fut.result()
            except Exception as exc:
                out[dom] = {
                    "domain": dom,
                    "verdict": "FAIL",
                    "error": str(exc)[:200],
                    "tags": [],
                    "verify_score": 0,
                }

    ok_n = sum(1 for r in out.values() if str(r.get("verdict", "")).startswith("PASS"))
    log.info(f"落地验证完成: {len(out)} 域 · PASS {ok_n}")
    return out


def write_excel(
    store: MemoryStore,
    excel_path: Path,
    *,
    meta: Optional[dict] = None,
    verify: bool = True,
    verify_limit: int = 40,
    verify_workers: int = 6,
) -> Path:
    """桌面 FB.xlsx：赚钱流程 SOP + 产品/竞品 + 验证 + Reddit 金帖。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    def _cell(v: Any) -> Any:
        if v is None:
            return ""
        if isinstance(v, bool):
            return "Y" if v else "N"
        if isinstance(v, (int, float)):
            return v
        return _clean_text(v)[:32000]

    excel_path = Path(excel_path).expanduser()
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    tables = build_money_tables(store)
    snap = store.snapshot()

    verify_map: dict[str, dict[str, Any]] = {}
    if verify:
        try:
            verify_map = verify_landings(store, limit=verify_limit, workers=verify_workers)
        except Exception:
            log.exception("落地验证失败，继续写表")
            verify_map = {}

    def _vget(domain: str) -> dict[str, Any]:
        d = (domain or "").lower().lstrip("www.")
        return verify_map.get(d) or {}

    try:
        from pipeline.reddit_fb_rules import (
            advice_from_scores,
            evaluate_money_pipeline,
            load_catalog,
            score_competitor_reddit,
        )
        reddit_catalog = load_catalog()
    except Exception:
        log.exception("加载 reddit_fb_rules 失败")
        reddit_catalog = {"posts": [], "rules": [], "post_count": 0, "rule_count": 0}

        def score_competitor_reddit(row, verify=None):  # type: ignore
            return 0, [], []

        def advice_from_scores(s, v=""):  # type: ignore
            return "观察"

        def evaluate_money_pipeline(row, verify=None, reddit_score=None):  # type: ignore
            return {
                "stage_label": "", "pipeline_rank": 0, "money_process": "",
                "s1": "", "s2": "", "s3": "", "s4": "", "s5": "", "s6": "",
                "next_action": "", "test_budget": "", "blockers": [], "price_hint": None,
            }

    # ── products ──
    prod_header = [
        "排名", "产品/落地", "域名", "品牌", "广告创意数", "覆盖词数", "主CTA", "最长投放天",
        "Reddit分", "流程阶段", "赚钱进程", "S1Meta", "S2LP", "S3客单", "S4履约", "S5交叉", "S6测门",
        "下一步", "测试预算", "Reddit旗标", "命中规则",
        "验证结论", "验证分", "标价$", "退货天", "Dropship", "落地页", "建议",
    ]
    prod_scored: list[tuple] = []
    for row in tables["products"][1:]:
        domain = str(row[2] or "")
        v = _vget(domain)
        r_in = {
            "domain": domain,
            "ads": int(row[4] or 0),
            "landings": 1,
            "keywords": int(row[5] or 0),
            "max_days": int(row[7] or 0),
            "top_cta": str(row[6] or ""),
            "sample_url": str(row[10] or ""),
            "title": str(row[1] or ""),
            "path": str(row[1] or ""),
        }
        rscore, flags, rules = score_competitor_reddit(r_in, v)
        try:
            old = int(row[8] or 0)
        except Exception:
            old = 0
        final = min(100, int(rscore * 0.75 + old * 0.25))
        pipe = evaluate_money_pipeline(r_in, v, reddit_score=final)
        verdict = str(v.get("verdict") or "")
        advice = advice_from_scores(final, verdict)
        if pipe.get("money_process") == "KILL" or verdict == "FAIL_DROPSHIP":
            advice = "跳过"
            final = min(final, 20)
        elif pipe.get("s6") == "READY":
            advice = "可测·小预算"
        rank_key = int(pipe.get("pipeline_rank") or 0) * 10 + final
        prod_scored.append((rank_key, final, pipe, r_in, v, flags, rules, advice, old, row))

    prod_scored.sort(key=lambda x: -x[0])
    prod_rows: list[list[Any]] = [prod_header]
    for i, (_rk, final, pipe, r_in, v, flags, rules, advice, old, row) in enumerate(prod_scored[:300], 1):
        prod_rows.append([
            i, row[1], r_in["domain"], row[3], row[4], row[5], row[6], row[7],
            final, pipe.get("stage_label") or "", pipe.get("money_process") or "",
            pipe.get("s1") or "", pipe.get("s2") or "", pipe.get("s3") or "",
            pipe.get("s4") or "", pipe.get("s5") or "", pipe.get("s6") or "",
            pipe.get("next_action") or "", pipe.get("test_budget") or "",
            ", ".join(flags), ",".join(rules),
            str(v.get("verdict") or "未测"),
            v.get("verify_score") if v else "",
            v.get("price_hint") if v else "",
            v.get("return_policy_days") if v else "",
            "Y" if v.get("has_dropship") else ("N" if v else ""),
            row[10], advice,
        ])
    tables["products"] = prod_rows

    # ── competitors ──
    comp_header = [
        "优先级", "域名", "品牌页", "广告数", "落地页数", "覆盖关键词数", "主CTA",
        "最长投放天", "最早投放", "粉丝",
        "Reddit分", "流程阶段", "赚钱进程", "S1Meta", "S2LP", "S3客单", "S4履约", "S5交叉", "S6测门",
        "下一步", "测试预算", "Reddit旗标", "命中规则",
        "验证结论", "验证分", "标价$", "退货天", "Dropship", "代表落地页", "建议",
    ]
    comp_scored: list[tuple] = []
    for row in tables["competitors"][1:]:
        domain = str(row[1] or "")
        v = _vget(domain)
        r_in = {
            "domain": domain,
            "ads": int(row[3] or 0),
            "landings": int(row[4] or 0),
            "keywords": int(row[5] or 0),
            "max_days": int(row[7] or 0),
            "top_cta": str(row[6] or ""),
            "sample_url": str(row[12] or ""),
            "title": str(row[2] or ""),
        }
        rscore, flags, rules = score_competitor_reddit(r_in, v)
        try:
            old = int(row[10] or 0)
        except Exception:
            old = 0
        final = min(100, int(rscore * 0.75 + old * 0.25))
        pipe = evaluate_money_pipeline(r_in, v, reddit_score=final)
        verdict = str(v.get("verdict") or "")
        advice = advice_from_scores(final, verdict)
        if pipe.get("money_process") == "KILL" or verdict == "FAIL_DROPSHIP":
            advice = "跳过"
            final = min(final, 20)
        elif pipe.get("s6") == "READY":
            advice = "可测·小预算"
        rank_key = int(pipe.get("pipeline_rank") or 0) * 10 + final
        comp_scored.append((rank_key, final, pipe, r_in, v, flags, rules, advice, old, row))

    comp_scored.sort(key=lambda x: -x[0])
    comp_rows: list[list[Any]] = [comp_header]
    for i, (_rk, final, pipe, r_in, v, flags, rules, advice, old, row) in enumerate(comp_scored, 1):
        comp_rows.append([
            i, domain if False else r_in["domain"], row[2], row[3], row[4], row[5], row[6],
            row[7], row[8], row[9],
            final, pipe.get("stage_label") or "", pipe.get("money_process") or "",
            pipe.get("s1") or "", pipe.get("s2") or "", pipe.get("s3") or "",
            pipe.get("s4") or "", pipe.get("s5") or "", pipe.get("s6") or "",
            pipe.get("next_action") or "", pipe.get("test_budget") or "",
            ", ".join(flags), ",".join(rules),
            str(v.get("verdict") or "未测"),
            v.get("verify_score") if v else "",
            v.get("price_hint") if v else "",
            v.get("return_policy_days") if v else "",
            "Y" if v.get("has_dropship") else ("N" if v else ""),
            row[12], advice,
        ])
    tables["competitors"] = comp_rows

    # ── pipeline mixed（含 funnel_score / intensity / priority_queue）──
    ppl_header = [
        "总排名", "类型", "域名", "名称", "Reddit分", "流程分", "漏斗分", "竞争强度", "赚钱进程",
        "优先级队列", "当前阶段",
        "S1Meta长跑创意", "S2LP_Offer", "S3客单带", "S4履约", "S5交叉", "S6测门",
        "卡点", "下一步动作", "测试预算", "创意数", "投放天", "标价$", "验证", "落地页",
    ]
    mixed: list[tuple[int, list[Any]]] = []

    def _mix_row(typ: str, final, pipe, r_in, v, name) -> list[Any]:
        return [
            0, typ, r_in.get("domain"), name,
            final, pipe.get("pipeline_rank"),
            pipe.get("funnel_score") if pipe.get("funnel_score") is not None else "",
            pipe.get("intensity") if pipe.get("intensity") is not None else "",
            pipe.get("money_process"), pipe.get("priority_queue"), pipe.get("stage_label"),
            pipe.get("s1"), pipe.get("s2"), pipe.get("s3"), pipe.get("s4"), pipe.get("s5"), pipe.get("s6"),
            "；".join(pipe.get("blockers") or [])[:140],
            pipe.get("next_action"), pipe.get("test_budget"),
            r_in.get("ads"), r_in.get("max_days"),
            pipe.get("price_hint") if pipe.get("price_hint") is not None else "",
            v.get("verdict") or "未测", r_in.get("sample_url"),
        ]

    for _rk, final, pipe, r_in, v, flags, rules, advice, old, row in prod_scored[:300]:
        name = r_in.get("title") or r_in.get("path") or ""
        mixed.append((int(pipe.get("pipeline_rank") or 0) * 1000 + final, _mix_row("产品", final, pipe, r_in, v, name)))
    for _rk, final, pipe, r_in, v, flags, rules, advice, old, row in comp_scored[:300]:
        mixed.append((int(pipe.get("pipeline_rank") or 0) * 1000 + final, _mix_row("竞品", final, pipe, r_in, v, r_in.get("title") or "")))

    # L1: READY + 优先级队列在前 → 用 (queue_priority, funnel) 排序
    _qprio = {"A-立刻可测": 0, "B-补交叉/客单": 1, "B-补Meta或交叉": 2, "C-补落地验证": 3, "C-观察": 4, "D-跳过": 5}
    mixed.sort(key=lambda kv: (
        _qprio.get(str(kv[1][9] if len(kv[1]) > 9 else ""), 9),
        -int(kv[1][6] if len(kv[1]) > 6 else 0),
        -kv[0],
    ))

    pipeline_rows: list[list[Any]] = [ppl_header]
    # L2: 同域名只留最靠前一条（产品优先）
    seen_dom: set[str] = set()
    _deduped: list[tuple[int, list[Any]]] = []
    stable = sorted(mixed, key=lambda kv: (0 if kv[1][1] == "产品" else 1, -kv[0]))
    for _k, row in stable:
        dom = str(row[2] or "").lower()
        if dom in seen_dom:
            continue
        seen_dom.add(dom)
        _deduped.append((_k, row))
    _deduped.sort(key=lambda kv: -kv[0])

    for i, (_k, row) in enumerate(_deduped[:250], 1):
        row[0] = i
        pipeline_rows.append(row)
    tables["pipeline"] = pipeline_rows



    # verify sheet
    verify_rows: list[list[Any]] = [[
        "域名", "品牌", "验证结论", "验证分", "HTTP", "错误",
        "LP标签", "技术分", "RedditLP分", "标价$", "Trust徽章",
        "单品LP", "折扣叠层", "退货天", "Dropship", "预信号分", "验证URL",
    ]]
    for dom, v in sorted(verify_map.items(), key=lambda kv: -int(kv[1].get("verify_score") or 0)):
        verify_rows.append([
            dom, v.get("page_name") or "", v.get("verdict") or "", v.get("verify_score") or 0,
            v.get("http_status") or 0, v.get("error") or "",
            ", ".join(v.get("tags") or []), v.get("bonus") or 0, v.get("reddit_bonus") or 0,
            v.get("price_hint") if v.get("price_hint") is not None else "",
            v.get("trust_badges") or 0,
            "Y" if v.get("is_single_product") else "N",
            v.get("discount_stack_count") or 0, v.get("return_policy_days") or 0,
            "Y" if v.get("has_dropship") else "N", v.get("pre_score") or 0, v.get("url") or "",
        ])

    rule_rows: list[list[Any]] = [["规则ID", "名称", "主题", "权重", "应用方式", "原话/要点", "源帖ID"]]
    for rule in reddit_catalog.get("rules") or []:
        rule_rows.append([
            rule.get("id"), rule.get("name"), rule.get("theme"), rule.get("weight"),
            rule.get("apply"), rule.get("quote"), ",".join(rule.get("sources") or []),
        ])

    gold_rows: list[list[Any]] = [["#", "subreddit", "score", "comments", "title", "url", "摘要"]]
    for i, p in enumerate((reddit_catalog.get("posts") or [])[:60], 1):
        gold_rows.append([
            i, p.get("subreddit"), p.get("score"), p.get("comments"),
            p.get("title"), p.get("url"), (p.get("selftext") or "")[:180],
        ])

    # SOP 说明页
    sop_rows = [
        ["步骤", "Reddit共识动作", "本表字段", "过关标准", "失败/卡点"],
        ["①", "Meta Ad Library 找长跑+多创意", "S1Meta / 最长投放天 / 广告创意数", "≥30天 且 ≥3创意（强：≥90天且≥5）", "单新广告/寿命短"],
        ["②", "点进独立站看 LP + Offer", "S2LP / 验证结论 / 落地页", "PASS + /products/ 或 Shopify；有保证/bundle更好", "Dropship三叠折扣、打不开"],
        ["③", "看客单带，毛利要另算COGS", "S3客单 / 标价$", "$20-89 冲动带；COGS 另用1688算3-5x", "无标价=UNKNOWN，不能说能赚钱"],
        ["④", "Boring/好运，避开易碎电子", "S4履约", "无 fragile/electronic 词", "重货易碎高退货"],
        ["⑤", "TT Creative Center + Amazon NR/BSR 交叉", "S5交叉", "本扫描固定 PENDING，需人工/下游交叉", "单FB证据不够"],
        ["⑥", "小预算测 → 再扩量", "S6测门 / 测试预算 / 下一步", "S1+S2过 + S4过 → READY", "前置关卡未过勿烧钱"],
        ["铁律", "高ROAS≠赚钱；先验证再扩；DRAFT不默认LIVE", "建议/赚钱进程", "TEST_READY 只代表可测", "禁止把信号当分润"],
    ]

    # ── 极致字段增强：富聚合 + Reddit赚钱案例 + 创意矩阵 ──
    try:
        from pipeline.extreme_table import (
            aggregate_domain_stats,
            copycat_clusters,
            creative_density_label,
            enrich_ad,
            fresh_page_risk,
            freshness_wave,
            hook_angle,
            longevity_label,
            match_money_cases,
            money_case_boost,
            offer_depth_grade,
            product_name_hint,
            reddit_thread_for_case,
            spend_confidence,
        )


        aid_kws: dict[str, set[str]] = defaultdict(set)
        for kw, aids in store.ads_by_keyword.items():
            for aid in aids:
                aid_kws[aid].add(kw)
        dom_stats = aggregate_domain_stats(store.ads_by_id, aid_kws)

        def _dom_extra(domain: str) -> dict[str, Any]:
            d = (domain or "").lower().lstrip("www.")
            return dom_stats.get(d) or {}

        # expand 赚钱产品
        ph = list(tables["products"][0]) + [
            "创意密度档", "寿命档", "活跃广告", "视频广告", "图文广告",
            "平台", "主格式", "Spend下限估", "Spend上限估", "曝光上限估",
            "Collation最大", "Offer广告数", "Hook广告数", "痛点广告数",
            "粉丝", "IG粉", "PageID数", "案例boost", "赚钱案例ID", "赚钱角度",
            "毛利待算提示", "案例源",
            "产品名(取自caption)", "Offer深度", "Copycat同款簇", "新页风险",
            "Spend置信", "新潮档", "钩子角度",
        ]
        # L3: copycat clusters built once
        _clusters = copycat_clusters(tables["products"])

        def _cluster_label(path: Any, clusters: dict) -> str:
            try:
                p = _re.sub(r"[/\s_-]+", "-", str(path or "").lower())[:60]
            except Exception:
                p = str(path or "").lower()
            c = clusters.get(p)
            if not c:
                return ""
            return f"{len(c['domains'])}域同款: {','.join(sorted(c['domains'])[:3])}"

        import re as _re

        new_prod = [ph]
        case_hits_rows: list[list[Any]] = [[
            "域名", "产品/品牌", "案例ID", "赚钱角度", "boost", "结构提示", "route", "来源", "强案例", "Reddit帖", "匹配文本摘要",
        ]]
        # L9: thread lookup fn
        def _thread(cid: str) -> str:
            try:
                return reddit_thread_for_case(cid)
            except Exception:
                return ""
        for row in tables["products"][1:]:
            domain = str(row[2] or "")
            st = _dom_extra(domain)
            text_blob = f"{row[1]} {row[3]} {row[26] if len(row) > 26 else ''} {st.get('sample_body') or ''}"
            hits = match_money_cases(text_blob, domain)
            boost, ids, angles = money_case_boost(hits)
            try:
                row = list(row)
                if len(row) > 8 and isinstance(row[8], (int, float)):
                    row[8] = min(100, int(row[8]) + boost)
            except Exception:
                row = list(row)
            price = None
            for j, name in enumerate(tables["products"][0]):
                if name in ("标价$", "标价提示") and j < len(row):
                    price = row[j]
                    break
            margin_hint = ""
            try:
                pf = float(price) if price not in (None, "") else None
            except Exception:
                pf = None
            if pf:
                margin_hint = f"售价~${pf:.0f} → 目标COGS≤${pf/3.5:.1f} (3.5x) / 理想≤${pf/5:.1f} (5x)；需1688实价"
            else:
                margin_hint = "缺标价：先探LP价再算3-5x毛利硬门"
            top_fmt = ""
            top_plat = ""
            if st:
                top_fmt = st["formats"].most_common(1)[0][0] if st.get("formats") else ""
                top_plat = ",".join(p for p, _ in st.get("platforms", Counter()).most_common(4))
            # round-2 enrichment
            phints = " / ".join(sorted(st.get("product_hints") or [])[:3])
            odep = int(st.get("offer_depth") or 0)
            ccl = _cluster_label(row[1], _clusters)
            frisk = fresh_page_risk(int(st.get("likes") or 0), int(st.get("max_days") or 0), int(st.get("ads") or 0))
            sconf = spend_confidence(st)
            fwave = freshness_wave(st)
            hangle = ",".join(f"{k}:{v}" for k, v in st.get("hook_angles", Counter()).most_common(3))
            new_prod.append(list(row) + [
                creative_density_label(int(st.get("ads") or row[4] or 0)),
                longevity_label(int(st.get("max_days") or row[7] or 0)),
                st.get("active_ads") or "",
                st.get("video_ads") or 0,
                st.get("image_ads") or 0,
                top_plat,
                top_fmt,
                st.get("spend_lower_sum") or 0,
                st.get("spend_upper_sum") or 0,
                st.get("imp_upper_sum") or 0,
                st.get("collation_max") or 1,
                st.get("offer_ads") or 0,
                st.get("hook_ads") or 0,
                st.get("pain_ads") or 0,
                st.get("likes") or 0,
                st.get("ig_followers") or 0,
                len(st.get("page_ids") or []),
                boost,
                ids,
                angles,
                margin_hint,
                " | ".join(str(h.get("source") or "") for h in hits[:2]),
                phints,
                f"{odep}类" if odep else "无",
                ccl,
                frisk,
                sconf,
                fwave,
                hangle,
            ])
            for h in hits[:5]:
                case_hits_rows.append([
                    domain, row[1], h.get("id"), h.get("money_angle"), h.get("boost"),
                    h.get("structure_hint"), h.get("route"), h.get("source"),
                    "Y" if h.get("strong") else "N",
                    str(text_blob)[:120],
                ])
        tables["products"] = new_prod

        # expand 竞争对手
        ch = list(tables["competitors"][0]) + [
            "创意密度档", "寿命档", "活跃广告", "视频广告", "图文广告", "平台", "主格式",
            "Spend下限估", "Spend上限估", "曝光上限估", "Collation最大",
            "Offer广告数", "Hook广告数", "痛点广告数", "IG粉", "PageID数",
            "案例boost", "赚钱案例ID", "赚钱角度", "毛利待算提示",
        ]
        new_comp = [ch]
        for row in tables["competitors"][1:]:
            domain = str(row[1] or "")
            st = _dom_extra(domain)
            text_blob = f"{row[2]} {domain} {st.get('sample_body') or ''}"
            hits = match_money_cases(text_blob, domain)
            boost, ids, angles = money_case_boost(hits)
            row = list(row)
            if len(row) > 10 and isinstance(row[10], (int, float)):
                row[10] = min(100, int(row[10]) + boost)
            top_fmt = st["formats"].most_common(1)[0][0] if st.get("formats") else ""
            top_plat = ",".join(p for p, _ in st.get("platforms", Counter()).most_common(4)) if st else ""
            price = None
            for j, name in enumerate(tables["competitors"][0]):
                if name in ("标价$", "标价提示") and j < len(row):
                    price = row[j]
                    break
            try:
                pf = float(price) if price not in (None, "") else None
            except Exception:
                pf = None
            if pf:
                margin_hint = f"售价~${pf:.0f} → COGS目标≤${pf/3.5:.1f}"
            else:
                margin_hint = "缺标价/COGS"
            ads_n = int(st.get("ads") or row[3] or 0)
            days_n = int(st.get("max_days") or row[7] or 0)
            new_comp.append(list(row) + [
                creative_density_label(ads_n), longevity_label(days_n),
                st.get("active_ads") or "", st.get("video_ads") or 0, st.get("image_ads") or 0,
                top_plat, top_fmt,
                st.get("spend_lower_sum") or 0, st.get("spend_upper_sum") or 0, st.get("imp_upper_sum") or 0,
                st.get("collation_max") or 1,
                st.get("offer_ads") or 0, st.get("hook_ads") or 0, st.get("pain_ads") or 0,
                st.get("ig_followers") or 0, len(st.get("page_ids") or []),
                boost, ids, angles, margin_hint,
            ])
            for h in hits[:3]:
                case_hits_rows.append([
                    domain, row[2], h.get("id"), h.get("money_angle"), h.get("boost"),
                    h.get("structure_hint"), h.get("route"), h.get("source"),
                    "Y" if h.get("strong") else "N", str(text_blob)[:120],
                ])
        tables["competitors"] = new_comp

        # 极致广告明细（L8 质量档）
        try:
            from pipeline.extreme_table import ad_quality_flag as _aqf

            def _qflag(ad: dict[str, Any]) -> str:
                try:
                    return _aqf(ad)
                except Exception:
                    return ""
        except Exception:
            _qflag = lambda ad: ""  # type: ignore
        ad_ext: list[list[Any]] = [[
            "ad_id", "关键词", "品牌", "page_id", "域名", "CTA", "投放天", "开始日", "结束日", "活跃",
            "格式", "平台", "视频", "图片", "collation",
            "spend_lo", "spend_hi", "imp_lo", "imp_hi", "currency",
            "粉丝", "IG", "IG粉", "Offer", "Hook", "痛点",
            "质量档", "信号分", "落地页", "caption", "文案摘要", "snapshot", "image_url", "video_url",
        ]]
        for aid, ad0 in store.ads_by_id.items():
            base = ad0 if "link_url" in ad0 else _ad_fields(ad0)
            ad = enrich_ad(base)
            sc, _ = _commerce_score(ad)
            ad_ext.append([
                aid,
                ", ".join(sorted(aid_kws.get(aid, set()))[:8]),
                ad.get("page_name") or "",
                ad.get("page_id") or "",
                ad.get("domain") or "",
                ad.get("cta") or "",
                ad.get("active_days") or 0,
                ad.get("start_date") or "",
                ad.get("end_date") or "",
                "Y" if ad.get("is_active") else "N",
                ad.get("display_format") or "",
                ad.get("platform_str") or "",
                "Y" if ad.get("has_video") else "N",
                "Y" if ad.get("has_image") else "N",
                ad.get("collation_count") or 1,
                ad.get("spend_lower") or 0,
                ad.get("spend_upper") or 0,
                ad.get("imp_lower") or 0,
                ad.get("imp_upper") or 0,
                ad.get("currency") or "",
                ad.get("page_like_count") or 0,
                ad.get("ig_handle") or "",
                ad.get("ig_followers") or 0,
                "Y" if ad.get("offer_flag") else "N",
                "Y" if ad.get("hook_flag") else "N",
                "Y" if ad.get("pain_flag") else "N",
                _qflag(ad),
                sc,
                ad.get("link_url") or "",
                ad.get("link_caption") or "",
                (ad.get("body") or ad.get("title") or "")[:200],
                ad.get("snapshot_url") or ad.get("ad_snapshot_url") or "",
                ad.get("image_url") or "",
                ad.get("video_url") or "",
            ])
        tables["ads"] = ad_ext

        # 创意矩阵：domain x format/platform
        creative_rows: list[list[Any]] = [[
            "域名", "品牌", "总广告", "密度档", "寿命档", "格式分布", "平台分布",
            "视频数", "图文数", "Offer数", "Hook数", "主落地",
        ]]
        for dom, st in sorted(dom_stats.items(), key=lambda kv: -int(kv[1].get("ads") or 0))[:200]:
            fmt = ", ".join(f"{k}:{v}" for k, v in st.get("formats", Counter()).most_common(5))
            plat = ", ".join(f"{k}:{v}" for k, v in st.get("platforms", Counter()).most_common(5))
            creative_rows.append([
                dom,
                " / ".join(sorted(st.get("page_names") or [])[:2]),
                st.get("ads") or 0,
                creative_density_label(int(st.get("ads") or 0)),
                longevity_label(int(st.get("max_days") or 0)),
                fmt, plat,
                st.get("video_ads") or 0, st.get("image_ads") or 0,
                st.get("offer_ads") or 0, st.get("hook_ads") or 0,
                st.get("sample_url") or "",
            ])

        tables["money_cases"] = case_hits_rows
        tables["creative_matrix"] = creative_rows
        extreme_ok = True
    except Exception:
        log.exception("极致字段增强失败，保留基础表")
        tables.setdefault("money_cases", [["域名", "说明"], ["", "增强失败"]])
        tables.setdefault("creative_matrix", [["域名", "说明"], ["", "增强失败"]])
        extreme_ok = False

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    good_fill = PatternFill("solid", fgColor="C6EFCE")
    bad_fill = PatternFill("solid", fgColor="FFC7CE")
    weak_fill = PatternFill("solid", fgColor="FFEB9C")

    focus_products = max(0, len(tables["products"]) - 1)
    focus_comp = max(0, len(tables["competitors"]) - 1)
    top_follow = 0
    try:
        advice_idx = tables["competitors"][0].index("建议")
        top_follow = sum(
            1 for r in tables["competitors"][1:]
            if str(r[advice_idx]) in ("重点跟", "可测·小预算", "可深挖/跟品验证")
        )
    except Exception:
        pass
    pass_n = sum(1 for v in verify_map.values() if str(v.get("verdict", "")).startswith("PASS"))
    fail_n = sum(1 for v in verify_map.values() if str(v.get("verdict", "")).startswith("FAIL"))
    ready_n = sum(1 for r in pipeline_rows[1:] if len(r) > 13 and r[13] == "READY")

    ws0 = wb.active
    ws0.title = "汇总"
    rows0 = [
        ("扫描时间", snap.get("started_at")),
        ("耗时秒", snap.get("elapsed_s")),
        ("选词方法", (meta or {}).get("method")),
        ("纯内存", "是"),
        ("HTTP服务", "否"),
        ("关键词成功", f"{snap.get('ok_count')}/{snap.get('done')}"),
        ("去重广告", snap.get("unique_ads")),
        ("竞争对手数", focus_comp),
        ("产品线索数", focus_products),
        ("重点跟/可测", top_follow),
        ("落地验证数", len(verify_map)),
        ("验证PASS", pass_n),
        ("验证FAIL", fail_n),
        ("流程READY可测", ready_n),
        ("Reddit金帖数", reddit_catalog.get("post_count") or len(reddit_catalog.get("posts") or [])),
        ("Reddit规则数", reddit_catalog.get("rule_count") or len(reddit_catalog.get("rules") or [])),
        ("赚钱案例命中行", max(0, len(tables.get("money_cases") or []) - 1)),
        ("创意矩阵域名", max(0, len(tables.get("creative_matrix") or []) - 1)),
        ("广告明细列数", len((tables.get("ads") or [[]])[0])),
        ("产品表列数", len((tables.get("products") or [[]])[0])),
        ("竞品表列数", len((tables.get("competitors") or [[]])[0])),
        ("本轮kw/h", snap["throughput"]["kw_per_hour"]),
        ("赚钱SOP", "①Meta长跑创意→②LP/Offer→③客单带→④履约→⑤TT/AMZ交叉→⑥小预算测"),
        ("说明", "先看赚钱流程READY + 赚钱案例；广告明细含spend/平台/视频；缺COGS不谈能赚钱"),
    ]

    ws0.append(["字段", "值"])
    for c in ws0[1]:
        c.fill = header_fill
        c.font = header_font
    for r in rows0:
        ws0.append([_cell(r[0]), _cell(r[1])])
    ws0.column_dimensions["A"].width = 16
    ws0.column_dimensions["B"].width = 78

    def _put(title: str, rows: list[list[Any]]) -> None:
        ws = wb.create_sheet(title)
        v_idx = a_idx = s6_idx = None
        if rows:
            try:
                v_idx = rows[0].index("验证结论")
            except ValueError:
                pass
            try:
                a_idx = rows[0].index("建议")
            except ValueError:
                pass
            try:
                s6_idx = rows[0].index("S6测门")
            except ValueError:
                try:
                    s6_idx = rows[0].index("S6测门") if False else None
                except Exception:
                    s6_idx = None
            if s6_idx is None:
                for j, name in enumerate(rows[0]):
                    if "S6" in str(name):
                        s6_idx = j
                        break
        for i, row in enumerate(rows):
            ws.append([_cell(x) for x in row])
            if i == 0:
                for c in ws[1]:
                    c.fill = header_fill
                    c.font = header_font
                continue
            fill = None
            if a_idx is not None and a_idx < len(row) and str(row[a_idx]) in ("重点跟", "可测·小预算", "可深挖/跟品验证"):
                fill = good_fill
            if s6_idx is not None and s6_idx < len(row) and str(row[s6_idx]) == "READY":
                fill = good_fill
            if v_idx is not None and v_idx < len(row):
                ver = str(row[v_idx])
                if ver.startswith("PASS"):
                    fill = good_fill
                elif ver.startswith("FAIL"):
                    fill = bad_fill
                elif ver == "WEAK":
                    fill = weak_fill
            if fill:
                for c in ws[i + 1]:
                    c.fill = fill
        for idx in range(1, min(27, (len(rows[0]) if rows else 1) + 1)):
            ws.column_dimensions[chr(64 + idx)].width = 14 if idx < 10 else 22

    _put("赚钱流程", pipeline_rows)
    _put("赚钱SOP", sop_rows)
    _put("赚钱产品", tables["products"])
    _put("竞争对手", tables["competitors"])
    _put("赚钱案例", tables.get("money_cases") or [["无数据"]])
    _put("创意矩阵", tables.get("creative_matrix") or [["无数据"]])
    _put("落地验证", verify_rows)
    _put("Reddit规则", rule_rows)
    _put("Reddit金帖", gold_rows)
    _put("广告明细", tables["ads"])


    ws_kw = wb.create_sheet("关键词")
    ws_kw.append(["关键词", "成功", "状态", "广告数", "耗时秒", "来源"])
    for c in ws_kw[1]:
        c.fill = header_fill
        c.font = header_font
    for r in store.results:
        ws_kw.append([
            _cell(r.keyword), "Y" if r.ok else "N", _cell(r.status),
            r.ads_count, r.duration_s, _cell(r.source),
        ])

    wb.save(excel_path)
    log.info(
        f"Excel 已写: {excel_path} | 流程READY {ready_n} | 竞品{focus_comp} 产品{focus_products} "
        f"验证{len(verify_map)} PASS{pass_n} 广告{snap.get('unique_ads')}"
    )
    return excel_path


def flush_store(store: MemoryStore, out_path: Path) -> Path:
    """把内存结果落盘（可选，不经 HTTP）。"""
    payload = store.snapshot()
    # 结果明细（不含完整 ads 正文时可仍保留计数）
    payload["results"] = [
        {
            "keyword": r.keyword,
            "ok": r.ok,
            "status": r.status,
            "ads_count": r.ads_count,
            "duration_s": r.duration_s,
            "source": r.source,
            "file": r.file,
            "error": r.error,
        }
        for r in store.results
    ]
    # 紧凑广告索引：只留关键字段，避免 GB 级 dump
    compact_ads = []
    for aid, ad in list(store.ads_by_id.items())[:50_000]:
        compact_ads.append(
            {
                "ad_archive_id": aid,
                "page_name": ad.get("page_name") or ad.get("pageName"),
                "link_url": ad.get("link_url") or ad.get("linkUrl") or ad.get("url"),
                "cta": ad.get("cta_type") or ad.get("cta_text") or ad.get("cta"),
                "domain": ad.get("domain"),
            }
        )
    payload["ads_sample"] = compact_ads
    payload["ads_sample_note"] = f"compact fields only; unique_ads={store.ads_total}"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return out_path


# ── CLI ─────────────────────────────────────────────────────────────────────


def main() -> int:
    parser = argparse.ArgumentParser(
        description="纯 Python 内存扫描 FB Ad Library（无 HTTP）"
    )
    src = parser.add_mutually_exclusive_group()
    src.add_argument("--keywords", help="逗号/换行分隔关键词")
    src.add_argument("--library", action="store_true", help="从词库/队列加载")
    src.add_argument(
        "--rabbit-hole",
        action="store_true",
        help="兔子洞选品法：队列 + 10 层 expand",
    )
    src.add_argument(
        "--bench",
        action="store_true",
        help="基准模式：小样本测速 + 日产能估算",
    )

    parser.add_argument("--limit", type=int, default=20, help="最多扫多少词")
    parser.add_argument("--workers", type=int, default=4, help="并发 worker")
    parser.add_argument("--country", default="US")
    parser.add_argument("--max-scrolls", type=int, default=5)
    parser.add_argument("--max-pages", type=int, default=0, help="0=用 FB_MAX_PAGES env")
    parser.add_argument("--headful", action="store_true", help="有头浏览器调试")
    parser.add_argument("--delay-sec", type=float, default=0.1)
    parser.add_argument("--hours", type=float, default=0.0, help=">0 时到点停")
    parser.add_argument("--no-keep-ads", action="store_true", help="内存不留广告正文，只计吞吐")
    parser.add_argument(
        "--disk-raw",
        action="store_true",
        help="写中间 raw JSON（默认关闭=纯内存）",
    )
    parser.add_argument(
        "--excel",
        default=str(Path.home() / "Desktop" / "FB.xlsx"),
        help="结束时写入的 Excel 路径（默认桌面 FB.xlsx）",
    )
    parser.add_argument("--no-excel", action="store_true", help="不写 Excel")
    parser.add_argument(
        "--verify",
        action="store_true",
        default=True,
        help="扫描后探测落地页验证数据（默认开）",
    )
    parser.add_argument("--no-verify", action="store_true", help="关闭落地页验证")
    parser.add_argument("--verify-limit", type=int, default=40, help="最多验证多少域名")
    parser.add_argument("--verify-workers", type=int, default=6, help="验证并发")
    parser.add_argument(
        "--out",
        default="",
        help="结果 JSON 路径；默认 logs/memory_scan_<ts>.json",
    )
    parser.add_argument("--no-flush", action="store_true", help="不写 JSON 汇总")
    parser.add_argument("--seeds", default="", help="兔子洞种子，逗号分隔")
    parser.add_argument("--log-level", default="INFO")
    parser.add_argument("--progress-every", type=int, default=5)
    args = parser.parse_args()

    setup_logging(level=args.log_level, log_file=str(LOG_DIR / "memory_scan.log"))
    os.environ.setdefault("FB_NO_HTTP", "1")

    limit = max(1, min(args.limit, 5000))
    select_meta: dict[str, Any] = {"method": "library"}
    pure_memory = not args.disk_raw

    if args.bench:
        limit = min(limit, 20)
        keywords = load_keywords_from_library(limit)
        select_meta = {"method": "bench/library"}
    elif args.keywords:
        keywords = _parse_text_items(args.keywords)[:limit]
        select_meta = {"method": "keywords"}
    elif args.rabbit_hole:
        seeds = _parse_text_items(args.seeds) if args.seeds else None
        keywords, select_meta = load_keywords_from_rabbit_hole(limit, seeds=seeds, expand=True)
    elif args.library:
        keywords = load_keywords_from_library(limit)
        select_meta = {"method": "library"}
    else:
        # 默认：兔子洞（用户要的选品路径）
        keywords, select_meta = load_keywords_from_rabbit_hole(limit, expand=True)

    if not keywords:
        print(json.dumps({"ok": False, "error": "关键词队列为空"}, ensure_ascii=False, indent=2))
        return 2

    print(
        f"[select] method={select_meta.get('method')} n={len(keywords)} "
        f"queue={select_meta.get('from_queue')} expand={select_meta.get('from_expand')} "
        f"pure_memory={pure_memory}",
        file=sys.stderr,
    )
    print(f"[select] sample={keywords[:8]}", file=sys.stderr)

    mp = args.max_pages if args.max_pages > 0 else None
    store = run_memory_scan(
        keywords,
        workers=max(1, args.workers),
        max_scrolls=max(1, min(args.max_scrolls, 20)),
        headless=not args.headful,
        country=args.country.strip().upper() or "US",
        max_pages=mp,
        delay_sec=max(0.0, args.delay_sec),
        keep_ads=not args.no_keep_ads,
        hours=max(0.0, args.hours),
        progress_every=max(1, args.progress_every),
        pure_memory=pure_memory,
    )

    snap = store.snapshot()
    snap["cli"] = {
        "bench": bool(args.bench),
        "rabbit_hole": bool(args.rabbit_hole) or select_meta.get("method") == "rabbit_hole",
        "library": bool(args.library),
        "pure_memory": pure_memory,
        "workers": args.workers,
        "limit": limit,
        "select_meta": select_meta,
    }

    excel_path = None
    if not args.no_excel:
        try:
            do_verify = (not args.no_verify) and bool(getattr(args, "verify", True))
            excel_path = write_excel(
                store,
                Path(args.excel),
                meta=select_meta,
                verify=do_verify,
                verify_limit=max(5, min(args.verify_limit, 120)),
                verify_workers=max(1, min(args.verify_workers, 12)),
            )
            snap["excel"] = str(excel_path)
            snap["verify"] = do_verify
        except Exception as exc:
            log.exception("写 Excel 失败")
            snap["excel_error"] = str(exc)
            print(f"[excel] 写入失败: {exc}", file=sys.stderr)


    out_path = None
    if not args.no_flush:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = Path(args.out) if args.out else (LOG_DIR / f"memory_scan_{ts}.json")
        flush_store(store, out_path)
        snap["flushed_to"] = str(out_path)

    print(json.dumps(snap, ensure_ascii=False, indent=2))

    daily = snap.get("daily_capacity") or {}
    print("\n=== 本轮扫描 ===", file=sys.stderr)
    print(
        f"词: ok={snap['ok_count']}/{snap['done']} · 去重广告={snap['unique_ads']} · "
        f"纯内存={pure_memory} · 兔子洞={snap['cli']['rabbit_hole']}",
        file=sys.stderr,
    )
    print(
        f"速率: {snap['throughput']['kw_per_hour']} kw/h · {snap['throughput']['ads_per_hour']} ads/h",
        file=sys.stderr,
    )
    print(
        f"可持续日产能: ~{daily.get('recommended_daily_keywords')} 词/天 "
        f"(sustainable {daily.get('sustainable_kph')} kw/h)",
        file=sys.stderr,
    )
    if excel_path:
        print(f"Excel: {excel_path}", file=sys.stderr)
    if out_path:
        print(f"JSON: {out_path}", file=sys.stderr)
    return 0 if store.ok_count or store.done else 1





if __name__ == "__main__":
    raise SystemExit(main())
